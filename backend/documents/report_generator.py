"""Generate PDF and Word documents from protocol data."""

from io import BytesIO
from datetime import datetime


def _protocol_to_plain_text(cond_protocols: dict, train_templates: dict, keys: list) -> str:
    """Build plain-text representation of selected protocols."""
    lines = []
    lines.append("AYUN Personal Protocol Plan")
    lines.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append("")
    lines.append("=" * 60)
    lines.append("")

    for ckey in keys:
        if ckey not in cond_protocols:
            continue
        p = cond_protocols[ckey]
        title = p.get("title", ckey)
        desc = p.get("description", "")
        lines.append(f"## {title}")
        if desc:
            lines.append(desc)
        lines.append("")

        tr = p.get("training", {}) or {}
        if tr:
            lines.append("Training")
            lines.append("-" * 40)
            for k in ["frequency", "intensity", "duration"]:
                v = tr.get(k, "")
                if v:
                    lines.append(f"  {k.title()}: {v}")
            focus = tr.get("focus", []) or []
            if focus:
                lines.append("  Focus:")
                for f in focus:
                    lines.append(f"    • {f}")
            lines.append("")

        nf = p.get("nutrition_focus", []) or []
        if nf:
            lines.append("Nutrition focus")
            lines.append("-" * 40)
            for n in nf:
                lines.append(f"  • {n}")
            lines.append("")

        recheck = p.get("recheck_interval", "")
        if recheck:
            lines.append(f"Recheck: {recheck}")
            lines.append("")

        rel = p.get("related_templates", []) or []
        if rel:
            lines.append("Training templates")
            lines.append("-" * 40)
            for tk in rel:
                if tk not in train_templates:
                    continue
                t = train_templates[tk]
                ttitle = t.get("title", tk)
                lines.append(f"  {ttitle}")
                for k in ["frequency", "intensity", "duration", "volume", "rest"]:
                    v = t.get(k, "")
                    if v:
                        lines.append(f"    {k.title()}: {v}")
                ex = t.get("exercises", t.get("modalities", [])) or []
                if ex:
                    for e in ex:
                        lines.append(f"    • {e}")
                lines.append("")
        lines.append("")

    return "\n".join(lines)


def protocol_to_pdf(
    cond_protocols: dict,
    train_templates: dict,
    keys: list,
) -> bytes:
    """Generate PDF bytes from selected protocol keys."""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
    except ImportError:
        raise ImportError("reportlab is required for PDF export. Run: pip install reportlab")

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    margin = 2 * cm
    y = height - margin
    line_height = 14

    def draw_text(s: str, font_size: int = 11, bold: bool = False):
        nonlocal y
        c.setFont("Helvetica-Bold" if bold else "Helvetica", font_size)
        # Simple wrap: split long lines
        max_chars = int((width - 2 * margin) / (font_size * 0.6))
        parts = []
        for line in s.split("\n"):
            while len(line) > max_chars:
                parts.append(line[:max_chars])
                line = line[max_chars:]
            parts.append(line)
        for part in parts:
            if y < margin + line_height:
                c.showPage()
                y = height - margin
            c.drawString(margin, y, part)
            y -= line_height

    text = _protocol_to_plain_text(cond_protocols, train_templates, keys)
    for line in text.split("\n"):
        if line.startswith("## "):
            if y < height - margin - 40:
                c.showPage()
                y = height - margin
            draw_text(line[3:], font_size=14, bold=True)
        elif line.startswith("=") or line.startswith("-"):
            y -= 4
        elif line.strip() == "":
            y -= line_height * 0.5
        else:
            draw_text(line)

    c.save()
    buf.seek(0)
    return buf.read()


def protocol_to_docx(
    cond_protocols: dict,
    train_templates: dict,
    keys: list,
) -> bytes:
    """Generate Word document bytes from selected protocol keys."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError("python-docx is required for Word export. Run: pip install python-docx")

    doc = Document()
    doc.add_heading("AYUN Personal Protocol Plan", 0)
    p = doc.add_paragraph()
    p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()

    for ckey in keys:
        if ckey not in cond_protocols:
            continue
        prot = cond_protocols[ckey]
        title = prot.get("title", ckey)
        desc = prot.get("description", "")
        doc.add_heading(title, level=1)
        if desc:
            doc.add_paragraph(desc)

        tr = prot.get("training", {}) or {}
        if tr:
            doc.add_heading("Training", level=2)
            for k in ["frequency", "intensity", "duration"]:
                v = tr.get(k, "")
                if v:
                    doc.add_paragraph(f"{k.title()}: {v}", style="List Bullet")
            focus = tr.get("focus", []) or []
            for f in focus:
                doc.add_paragraph(f, style="List Bullet 2")

        nf = prot.get("nutrition_focus", []) or []
        if nf:
            doc.add_heading("Nutrition focus", level=2)
            for n in nf:
                doc.add_paragraph(n, style="List Bullet")

        recheck = prot.get("recheck_interval", "")
        if recheck:
            doc.add_paragraph(f"Recheck: {recheck}")

        rel = prot.get("related_templates", []) or []
        if rel:
            doc.add_heading("Training templates", level=2)
            for tk in rel:
                if tk not in train_templates:
                    continue
                t = train_templates[tk]
                ttitle = t.get("title", tk)
                doc.add_heading(ttitle, level=3)
                for k in ["frequency", "intensity", "duration", "volume", "rest"]:
                    v = t.get(k, "")
                    if v:
                        doc.add_paragraph(f"{k.title()}: {v}", style="List Bullet")
                ex = t.get("exercises", t.get("modalities", [])) or []
                for e in ex:
                    doc.add_paragraph(e, style="List Bullet 2")

        doc.add_paragraph()

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def doctor_review_to_pdf(export_data: dict) -> bytes:
    """Generate PDF bytes from Doctor Review export data."""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
    except ImportError:
        raise ImportError("reportlab is required for PDF export. Run: pip install reportlab")

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    margin = 2 * cm
    y = height - margin
    line_height = 14

    def draw_text(s: str, font_size: int = 11, bold: bool = False):
        nonlocal y
        c.setFont("Helvetica-Bold" if bold else "Helvetica", font_size)
        max_chars = int((width - 2 * margin) / (font_size * 0.6))
        parts = []
        for line in (s or "").replace("\r", "").split("\n"):
            while len(line) > max_chars:
                parts.append(line[:max_chars])
                line = line[max_chars:]
            parts.append(line)
        for part in parts:
            if y < margin + line_height:
                c.showPage()
                y = height - margin
            c.drawString(margin, y, part)
            y -= line_height

    meta = export_data.get("_meta", {})
    member = export_data.get("member", {})
    recs = export_data.get("recommendations", [])

    draw_text("AYUN Doctor Review Plan", font_size=16, bold=True)
    draw_text(f"Review date: {meta.get('review_date', '')}")
    draw_text(f"Reviewed by: {meta.get('reviewed_by', '')}")
    y -= line_height

    m_name = member.get("name") or "N/A"
    m_grp = member.get("group", "")
    draw_text(f"Member: {m_name} | {m_grp}", font_size=10)
    draw_text(f"Age: {member.get('age')} | Sex: {member.get('sex')}")
    y -= line_height

    draw_text("Recommendations", font_size=12, bold=True)
    y -= 4

    for rec in recs:
        prio = rec.get("reviewed_priority") or rec.get("original_priority", "")
        title = rec.get("title", "")
        draw_text(f"[{prio}] {title}", font_size=11, bold=True)
        items = rec.get("action_items", "")
        if items:
            for line in items.strip().split("\n"):
                txt = line.strip().lstrip("• ").strip()
                if txt:
                    draw_text(f"  • {txt}", font_size=10)
        notes = (rec.get("clinician_notes") or "").strip()
        if notes:
            draw_text(f"  Clinician notes: {notes}", font_size=9)
        y -= line_height

    c.save()
    buf.seek(0)
    return buf.read()
