"""
Threshold test suite: every AYUN cut-point, every sex x age group (14 groups),
values just below / at / just above. Oracle = AYUN's own tables (workbook,
visceral-fat note, One-Year Outcomes Overview), not the app's YAML.
Checks (1) rate_value() rating and (2) whether the full pipeline
(get_recommendations) shows an action card for that marker.
"""
import sys, json, re
sys.path[:0] = [sys.argv[2], "/Users/jiaqiyu/Desktop/AYUN/backend", "/Users/jiaqiyu/Desktop/AYUN"]
from ayun_bands import load, VAT, APP_GROUPS, parse_cell
from rule_engine import rate_value, ACTION_RATINGS, _PARAM_MAP
from recommendation_api import get_recommendations

PARAM_KEYS = {  # AYUN test -> app member_data key (right side where L/R)
    "vo2max": "Physiological::Vo2Max", "grip_strength": "Physiological::Gripstrength Rechts",
    "counter_movement_jump": "Physiological::Counter movement jump", "push_up_30s": "Physiological::Push Up 30s",
    "y_balance_composite": "Physiological::Y-Balance Composite Score",
    "shoulder_flexion": "Physiological::Shoulder Flexion Right", "hip_extension": "Physiological::Hip Extension Right",
    "hip_flexion": "Physiological::Hip Flexion Right", "knee_extension": "Physiological::Knee Extension Rechts",
    "knee_flexion": "Physiological::Knee Flexion Rechts", "elbow_extension": "Physiological::Triceps Strength Rechts",
    "elbow_flexion": "Physiological::Biceps Strength Rechts",
}
FIVE = {"below average": "poor", "fair": "below_average", "good": "average", "very good": "good", "excellent": "excellent"}
FOUR = {"below average": "poor", "fair": "below_average", "good": "average", "very good": "excellent"}
ROM = {"shoulder_flexion", "hip_extension", "hip_flexion"}
AGE_OF = {"age_18_29": 25, "age_30_39": 35, "age_40_49": 45, "age_50_59": 55, "age_60_69": 65, "age_70_79": 75, "age_80_99": 85}
AGES = {"age_18_29": (20, 29), "age_30_39": (30, 39), "age_40_49": (40, 49), "age_50_59": (50, 59),
        "age_60_69": (60, 69), "age_70_79": (70, 79), "age_80_99": (80, 99)}
from ayun_bands import pick

def ayun_tiers(t):
    """[(name, lo, hi)] from the parsed AYUN cell."""
    out = []
    for name, (sign, a, b) in t.items():
        if name == "below average": out.append((name, -1e9, a))
        elif name == "excellent" or (name == "very good" and b is None): out.append((name, a, 1e9))
        else: out.append((name, a, b if b is not None else 1e9))
    return out

def expected(tiers, v, mapping):
    hits = [mapping[n] for n, lo, hi in tiers if lo - 1e-9 <= v <= hi + 1e-9]
    return sorted(set(hits))

results = []
data = load(sys.argv[1])
import openpyxl
_ws = openpyxl.load_workbook(sys.argv[1], data_only=True).worksheets[0]
_rows = {str(r[0]).strip().lower(): r for r in _ws.iter_rows(values_only=True) if r[0]}
_LABEL = {"shoulder_flexion": "shoulder flexion right", "hip_extension": "hip extension right", "hip_flexion": "hip flexion right",
          "knee_extension": "knee extension right (seated)", "knee_flexion": "knee flexion right (prone)",
          "elbow_extension": "elbow extension right", "elbow_flexion": "elbow flexion right", "grip_strength": "hand grip right",
          "vo2max": "vo2max", "counter_movement_jump": "cmj", "push_up_30s": "30 sec push up", "y_balance_composite": "y balance total"}
data_rows = {k: _rows[v] for k, v in _LABEL.items()}
for param, sexes in data.items():
    mapping = FOUR if param in ROM else FIVE
    key = PARAM_KEYS[param]
    for sex, groups in sexes.items():
        for grp, (lo_map, (col_groups, _), _t) in groups.items():
          for age in AGES[grp]:
            row = data_rows[param]
            if param in ("knee_extension", "knee_flexion", "elbow_extension", "elbow_flexion") and sex == "male":
                col_groups = [(0, 49, col_groups[0][2]), (50, 59, col_groups[1][2]), (60, 120, col_groups[2][2])]  # AYUN overlaps at 50
            tiers = parse_cell(row[pick(col_groups, age)])
            tl = ayun_tiers(tiers)
            cuts = sorted({lo for n, lo, hi in tl if n != "below average"})
            step = 1 if param == "push_up_30s" else 0.1
            for c in cuts:
                for v in (round(c - step, 2), round(c, 2), round(c + step, 2)):
                    exp = expected(tl, v, mapping)
                    got = rate_value(key, v, sex.capitalize(), age).rating
                    if not exp:
                        verdict = "gap"  # value falls between AYUN rows (e.g. 41.95)
                    else:
                        verdict = "pass" if got in exp else "fail"
                    rec = get_recommendations({"Demographic::Age": age, "Demographic::Sex": sex.capitalize(), key: v})
                    card = any(key in t.parameters for t in rec.triggered)
                    want_card = bool(exp) and all(e in ACTION_RATINGS for e in exp)
                    no_card = bool(exp) and all(e not in ACTION_RATINGS for e in exp)
                    card_ok = (card if want_card else (not card) if no_card else True)
                    results.append(dict(test=param, sex=sex, group=grp, age=age, cut=c, value=v,
                                        expected=exp, got=got, rating=verdict, card=card, card_ok=card_ok,
                                        ambiguous=len(exp) > 1))

# Visceral fat (AYUN note): tiers excellent / very good / good / fair / bad by upper bounds
VAT_MAP = ["excellent", "good", "average", "elevated", "high"]
for sex, groups in VAT.items():
    for grp, ups in groups.items():
        age = AGE_OF[grp]
        bounds = [(0, ups[0])] + [(ups[i] + 1, ups[i + 1]) for i in range(3)] + [(ups[3] + 1, 1e9)]
        for i in range(4):
            c = ups[i] + 1  # first value of the next (worse) tier
            for v in (c - 1, c, c + 1):
                exp = [VAT_MAP[j] for j, (lo, hi) in enumerate(bounds) if lo <= v <= hi]
                if (sex, grp, v) in {("female", "age_18_29", 344), ("female", "age_70_79", 1538), ("female", "age_80_99", 1538)}:
                    exp = ["elevated", "high"]  # AYUN note lists this value in two tiers
                got = rate_value("Body Composition::Visceral Fat Index (raw value from DEXA)", v, sex.capitalize(), age).rating
                rec = get_recommendations({"Demographic::Age": age, "Demographic::Sex": sex.capitalize(),
                                           "Body Composition::Visceral Fat Index (raw value from DEXA)": v})
                card = any("Visceral" in p for t in rec.triggered for p in t.parameters)
                want = exp[0] in ACTION_RATINGS
                results.append(dict(test="visceral_fat", sex=sex, group=grp, age=age, cut=c, value=v, expected=exp, got=got,
                                    rating="pass" if got in exp else "fail", card=card, card_ok=(card == want), ambiguous=False))

# Blood cut-offs (AYUN One-Year Outcomes Overview): at/above cut -> action
BLOOD = [("Blood Biomarkers::LDL", 3.0, 0.01), ("Blood Biomarkers::HbA1c", 5.7, 0.01), ("Blood Biomarkers::HOMA Index", 2.5, 0.01),
         ("Blood Biomarkers::Triglycerides", 1.7, 0.01), ("Blood Biomarkers::hsCRP", 3.0, 0.01)]
for sex in ("male", "female"):
    for grp, age in AGE_OF.items():
        for key, c, d in BLOOD:
            for v in (round(c - d, 2), c, round(c + d, 2)):
                got = rate_value(key, v, sex.capitalize(), age).rating
                want_action = v >= c
                rec = get_recommendations({"Demographic::Age": age, "Demographic::Sex": sex.capitalize(), key: v})
                card = any(key in t.parameters for t in rec.triggered)
                results.append(dict(test=key.split("::")[1], sex=sex, group=grp, age=age, cut=c, value=v,
                                    expected=["action" if want_action else "no action"], got=got,
                                    rating="pass" if (got in ACTION_RATINGS) == want_action else "fail",
                                    card=card, card_ok=(card == want_action), ambiguous=False))

# ── Combined profiles: many markers at once, plus blank / zero / missing demographics ──
import random
rng = random.Random(42)
by_group = {}
for r in results:
    if r["rating"] == "pass" and not r["ambiguous"] and r["test"] in PARAM_KEYS:
        by_group.setdefault((r["sex"], r["age"]), {}).setdefault(r["test"], []).append(r)
combined = []
for (sex, age), tests in sorted(by_group.items()):
    for variant in ("full", "one_zero", "one_blank", "no_sex", "no_age"):
        for rep_i in range(3):
            picks = {t: rng.choice(rows) for t, rows in tests.items()}
            member = {"Demographic::Age": age, "Demographic::Sex": sex.capitalize()}
            for t, r in picks.items():
                member[PARAM_KEYS[t]] = r["value"]
            dropped = None
            if variant in ("one_zero", "one_blank"):
                dropped = rng.choice(sorted(picks))
                member[PARAM_KEYS[dropped]] = 0 if variant == "one_zero" else None
            if variant == "no_sex": member.pop("Demographic::Sex")
            if variant == "no_age": member.pop("Demographic::Age")
            want = set()
            if variant not in ("no_sex", "no_age"):
                want = {PARAM_KEYS[t] for t, r in picks.items() if t != dropped and r["expected"][0] in ACTION_RATINGS}
            rec = get_recommendations(member)
            fired = {p for t in rec.triggered for p in t.parameters if p in {PARAM_KEYS[t] for t in picks}}
            asks = any(t.category == "Missing data" for t in rec.triggered)
            ok = fired == want and asks == (variant in ("no_sex", "no_age"))
            combined.append(dict(sex=sex, age=age, variant=variant, markers=len(picks), want=sorted(want), fired=sorted(fired),
                                 asks_for_missing=asks, ok=ok))
json.dump({"single": results, "combined": combined}, open(sys.argv[3], "w"), indent=0)
from collections import Counter as C2
print("combined:", len(combined), "pass:", sum(c["ok"] for c in combined), dict(C2((c["variant"], c["ok"]) for c in combined)))
for c in [c for c in combined if not c["ok"]][:10]: print("CFAIL", c)
from collections import Counter
tested = [r for r in results if r["rating"] != "gap"]
print("cases:", len(results), "| tested:", len(tested), "| gaps:", len(results) - len(tested))
print("rating:", Counter(r["rating"] for r in tested))
print("card:", Counter(r["card_ok"] for r in tested))
print("groups:", len({(r['sex'], r['group']) for r in results}), "tests:", sorted({r['test'] for r in results}))
for r in [r for r in tested if r["rating"] == "fail" or not r["card_ok"]][:40]:
    print("FAIL", r["test"], r["sex"], r["group"], "v=", r["value"], "exp", r["expected"], "got", r["got"], "card", r["card"], "card_ok", r["card_ok"])
"""Parse AYUN reference tables into {param: {sex: {app_age_group: {tier_lo...}}}}."""
import re, sys, openpyxl

TIER_RE = re.compile(r"(below average|fair|good|very good|excellent)\s*:?\s*(?:≤|≥|<|>)?\s*(\d+(?:[.,]\d+)?)\s*%?\s*(?:-\s*(\d+(?:[.,]\d+)?))?", re.I)
APP_GROUPS = [("age_18_29",20),("age_30_39",30),("age_40_49",40),("age_50_59",50),("age_60_69",60),("age_70_79",70),("age_80_99",80)]

def parse_cell(text):
    out = {}
    t = str(text).replace("dgrees","").replace("kg","")
    # handle "very good" before "good"
    for m in re.finditer(r"(below average|very good|fair|good|excellent)\s*:?\s*(≤|≥)?\s*(\d+(?:[.,]\d+)?)\s*%?\s*(?:-\s*(\d+(?:[.,]\d+)?))?", t, re.I):
        name = m.group(1).lower()
        if name in out: continue
        a = float(m.group(3).replace(",", ".")); b = m.group(4)
        out[name] = (m.group(2), a, float(b.replace(",", ".")) if b else None)
    return out

def lows(tiers, four_tier=False):
    """Lower bound of each tier above 'below average' in app naming."""
    names = ["fair","good","very good"] + ([] if four_tier else ["excellent"])
    app = ["below_average","average","good","excellent"] if not four_tier else ["below_average","average","excellent"]
    res = {}
    for n, a in zip(names, app):
        if n not in tiers: return None
        sign, x, y = tiers[n]
        res[a] = x
    return res

def pick(groups, age):
    """groups: list of (lo, hi, col). Return col for age."""
    for lo, hi, col in groups:
        if lo <= age <= hi: return col
    return groups[-1][2] if age > groups[-1][1] else groups[0][2]

def load(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    rows = {str(r[0]).strip().lower(): r for r in ws.iter_rows(values_only=True) if r[0]}
    spec = {
        # param: (row label, four_tier, women groups, men groups) groups=(lo,hi,colIndex)
        "shoulder_flexion": ("shoulder flexion right", True, [(0,44,3),(45,120,4)], [(0,44,5),(45,120,6)]),
        "hip_extension": ("hip extension right", True, [(0,44,3),(45,120,4)], [(0,44,5),(45,120,6)]),
        "hip_flexion": ("hip flexion right", True, [(0,44,3),(45,120,4)], [(0,44,5),(45,120,6)]),
        "knee_extension": ("knee extension right (seated)", False, [(0,49,3),(50,59,4),(60,120,5)], [(0,49,6),(50,59,7),(60,120,8)]),
        "knee_flexion": ("knee flexion right (prone)", False, [(0,49,3),(50,59,4),(60,120,5)], [(0,49,6),(50,59,7),(60,120,8)]),
        "elbow_extension": ("elbow extension right", False, [(0,49,3),(50,59,4),(60,120,5)], [(0,49,6),(50,59,7),(60,120,8)]),
        "elbow_flexion": ("elbow flexion right", False, [(0,49,3),(50,59,4),(60,120,5)], [(0,49,6),(50,59,7),(60,120,8)]),
        "grip_strength": ("hand grip right", False, [(0,29,3),(30,39,4),(40,49,5),(50,59,6),(60,69,7),(70,120,8)], [(0,29,9),(30,39,10),(40,49,11),(50,59,12),(60,69,13),(70,120,14)]),
        "vo2max": ("vo2max", False, [(0,29,3),(30,39,4),(40,49,5),(50,59,6),(60,69,7),(70,120,8)], [(0,29,9),(30,39,10),(40,49,11),(50,59,12),(60,69,13),(70,120,14)]),
        "counter_movement_jump": ("cmj", False, [(0,29,3),(30,39,4),(40,49,5),(50,59,6),(60,120,7)], [(0,29,8),(30,39,9),(40,49,10),(50,59,11),(60,120,12)]),
        "push_up_30s": ("30 sec push up", False, [(0,29,3),(30,39,4),(40,49,5),(50,59,6),(60,120,7)], [(0,29,8),(30,39,9),(40,49,10),(50,59,11),(60,120,12)]),
        "y_balance_composite": ("y balance total", False, [(0,39,3),(40,49,4),(50,59,5),(60,69,6),(70,120,7)], [(0,39,8),(40,49,9),(50,59,10),(60,69,11),(70,120,12)]),
    }
    out = {}
    for param, (label, four, wg, mg) in spec.items():
        row = rows[label]
        for sex, groups in (("female", wg), ("male", mg)):
            for grp, age in APP_GROUPS:
                col = pick(groups, age)
                tiers = parse_cell(row[col])
                lo = lows(tiers, four)
                out.setdefault(param, {}).setdefault(sex, {})[grp] = (lo, (groups, col), tiers)
    return out

# Visceral fat (AYUN 'Visceral fat info' note). Typos 2030->1030 (male 50-59) and 2055->1055 (female 70-99) corrected.
VAT = {
  "female": {"age_18_29":[30,97,163,344], "age_30_39":[53,157,260,515], "age_40_49":[104,258,414,752],
             "age_50_59":[203,424,645,1061], "age_60_69":[331,608,886,1356], "age_70_79":[440,747,1054,1538], "age_80_99":[440,747,1054,1538]},
  "male":   {"age_18_29":[59,183,308,579], "age_30_39":[192,440,689,1126], "age_40_49":[370,728,1087,1648],
             "age_50_59":[581,1029,1477,2117], "age_60_69":[803,1296,1789,2443], "age_70_79":[981,1508,2036,2698], "age_80_99":[981,1508,2036,2698]},
}
